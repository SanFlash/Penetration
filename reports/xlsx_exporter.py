"""Create a shareable XLSX penetration assessment report with embedded evidence."""
from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SEVERITY_FILL = {
    "Critical": "FF4D6D",
    "High": "FF8A4C",
    "Medium": "FFC857",
    "Low": "38E8A0",
    "Info": "7D93AD",
}


def _safe_image(path: str | None) -> str | None:
    if not path:
        return None
    p = Path(str(path))
    return str(p) if p.is_file() and p.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"} else None


def _excel_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _find_image(finding: dict, evidence_gallery: list[dict]) -> str | None:
    candidates = []
    candidates.extend(finding.get("screenshots") or [])
    if finding.get("screenshot"):
        candidates.append(finding["screenshot"])
    candidates.extend(
        x.get("path") for x in evidence_gallery
        if x.get("name") and (
            x.get("name") in {os.path.basename(str(p)) for p in candidates}
        )
    )
    for candidate in candidates:
        image = _safe_image(candidate)
        if image:
            return image
    return None


def generate_xlsx(report: dict, output_path: str) -> str:
    """Generate a professional XLSX report and embed evidence images in Findings."""
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Summary"
    findings_ws = wb.create_sheet("Failed Issues")
    evidence_ws = wb.create_sheet("Evidence")
    remediation_ws = wb.create_sheet("Remediation")
    coverage_ws = wb.create_sheet("Coverage")

    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="102A43")
    sub_fill = PatternFill("solid", fgColor="EAF2F8")

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False

    # Executive Summary
    ws["A1"] = "SENTINEL — PENETRATION ASSESSMENT REPORT"
    ws["A1"].font = Font(size=20, bold=True, color="FFFFFF")
    ws["A1"].fill = header_fill
    ws.merge_cells("A1:F1")
    ws["A3"] = "Target"
    ws["B3"] = _excel_text(report.get("target"))
    ws["A4"] = "Generated"
    ws["B4"] = _excel_text(report.get("generated_at"))
    ws["A5"] = "Profile"
    ws["B5"] = _excel_text((report.get("metadata") or {}).get("profile", "-"))
    metrics = [
        ("Unique Findings", report.get("unique_findings", 0)),
        ("Raw Observations", report.get("raw_findings", 0)),
        ("Total Observations", report.get("total_observations", 0)),
        ("Evidence Images", len(report.get("evidence_gallery") or [])),
    ]
    for idx, (label, value) in enumerate(metrics, start=7):
        ws.cell(idx, 1, label).font = Font(bold=True)
        ws.cell(idx, 1).fill = sub_fill
        ws.cell(idx, 2, value)
        ws.cell(idx, 1).border = ws.cell(idx, 2).border = border

    ws["D3"] = "Severity"
    ws["E3"] = "Count"
    for cell in ws[3][3:5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for row, severity in enumerate(("Critical", "High", "Medium", "Low", "Info"), start=4):
        ws.cell(row, 4, severity)
        ws.cell(row, 5, (report.get("severity_summary") or {}).get(severity, 0))
        ws.cell(row, 4).fill = PatternFill("solid", fgColor=SEVERITY_FILL[severity])
        ws.cell(row, 4).font = Font(bold=True)
        ws.cell(row, 4).border = ws.cell(row, 5).border = border

    # Failed Issues — one row per reported finding, with embedded evidence image.
    headers = [
        "Issue ID", "Failed Issue / Finding", "Severity", "Confidence", "Category",
        "Method", "URL", "Parameter", "Evidence / Proof", "Impact", "Remediation", "Evidence Image"
    ]
    for col, header in enumerate(headers, 1):
        cell = findings_ws.cell(1, col, header)
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    findings_ws.freeze_panes = "A2"
    findings_ws.auto_filter.ref = f"A1:L{max(1, len(report.get('findings') or []) + 1)}"

    for row_idx, finding in enumerate(report.get("findings") or [], start=2):
        values = [
            finding.get("id"),
            finding.get("title"),
            finding.get("severity"),
            finding.get("confidence"),
            finding.get("category"),
            finding.get("method", "GET"),
            finding.get("url"),
            finding.get("parameter"),
            finding.get("evidence"),
            finding.get("impact"),
            finding.get("remediation"),
            "Attached below",
        ]
        for col, value in enumerate(values, 1):
            cell = findings_ws.cell(row_idx, col, _excel_text(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        sev = _excel_text(finding.get("severity"))
        if sev in SEVERITY_FILL:
            findings_ws.cell(row_idx, 3).fill = PatternFill("solid", fgColor=SEVERITY_FILL[sev])

        image_path = _find_image(finding, report.get("evidence_gallery") or [])
        if image_path:
            try:
                img = XLImage(image_path)
                original_width = max(1, img.width)
                original_height = max(1, img.height)
                # Fit the screenshot inside a predictable Excel cell area while
                # preserving its original aspect ratio.
                ratio = min(480 / original_width, 300 / original_height, 1)
                img.width = max(1, int(original_width * ratio))
                img.height = max(1, int(original_height * ratio))
                findings_ws.add_image(img, f"L{row_idx}")
                findings_ws.cell(row_idx, 12, "Embedded evidence image")
                findings_ws.row_dimensions[row_idx].height = 225
            except Exception as exc:
                findings_ws.cell(row_idx, 12, f"Image unavailable: {exc}")

    widths = [16, 34, 12, 12, 24, 10, 48, 20, 58, 45, 48, 34]
    for idx, width in enumerate(widths, 1):
        findings_ws.column_dimensions[get_column_letter(idx)].width = width

    # Evidence inventory.
    evidence_headers = ["Evidence", "Type", "Size (KB)", "Relative Path", "Finding IDs"]
    for col, header in enumerate(evidence_headers, 1):
        c = evidence_ws.cell(1, col, header)
        c.fill = header_fill
        c.font = Font(bold=True, color="FFFFFF")
        c.border = border
    for idx, item in enumerate(report.get("evidence_gallery") or [], start=2):
        vals = [item.get("name"), item.get("kind"), round((item.get("size") or 0) / 1024, 1), item.get("path"), ""]
        for col, value in enumerate(vals, 1):
            evidence_ws.cell(idx, col, _excel_text(value)).alignment = Alignment(vertical="top", wrap_text=True)
            evidence_ws.cell(idx, col).border = border
    evidence_ws.freeze_panes = "A2"

    # Remediation.
    remediation = report.get("remediation") or {}
    rh = ["Priority", "Severity", "Category", "Action", "Impact", "How to Validate", "Affected URLs"]
    for col, header in enumerate(rh, 1):
        c = remediation_ws.cell(1, col, header)
        c.fill = header_fill
        c.font = Font(bold=True, color="FFFFFF")
        c.border = border
    for idx, action in enumerate(remediation.get("actions") or [], start=2):
        vals = [
            action.get("priority"), action.get("severity"), action.get("category"),
            action.get("fix") or action.get("title"), action.get("impact"),
            action.get("validation"), action.get("affected_urls"),
        ]
        for col, value in enumerate(vals, 1):
            remediation_ws.cell(idx, col, _excel_text(value)).alignment = Alignment(vertical="top", wrap_text=True)
            remediation_ws.cell(idx, col).border = border
    remediation_ws.freeze_panes = "A2"

    # Browser/viewport coverage.
    ch = ["Browser", "Viewport", "URL", "Status", "Load (ms)", "Console Errors", "Network Failures", "Horizontal Overflow", "Screenshot"]
    for col, header in enumerate(ch, 1):
        c = coverage_ws.cell(1, col, header)
        c.fill = header_fill
        c.font = Font(bold=True, color="FFFFFF")
        c.border = border
    coverage = ((report.get("metadata") or {}).get("ui_responsive")
                or (report.get("metadata") or {}).get("compatibility") or {})
    for idx, item in enumerate(coverage.get("results") or [], start=2):
        vals = [
            item.get("browser"), item.get("viewport"), item.get("url"), item.get("status"),
            item.get("load_ms"), len(item.get("console_errors") or []),
            len(item.get("request_failures") or []),
            "YES" if item.get("horizontal_overflow") else "NO",
            item.get("screenshot"),
        ]
        for col, value in enumerate(vals, 1):
            coverage_ws.cell(idx, col, _excel_text(value)).alignment = Alignment(vertical="top", wrap_text=True)
            coverage_ws.cell(idx, col).border = border
    coverage_ws.freeze_panes = "A2"

    for sheet in wb.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.5
        sheet.page_margins.bottom = 0.5

    wb.save(out)
    return str(out)


def validate_xlsx(path: str) -> bool:
    """Open the generated workbook to verify it is a valid XLSX package."""
    wb = load_workbook(path, read_only=False)
    try:
        required = {"Executive Summary", "Failed Issues", "Evidence", "Remediation", "Coverage"}
        return required.issubset(set(wb.sheetnames))
    finally:
        wb.close()
